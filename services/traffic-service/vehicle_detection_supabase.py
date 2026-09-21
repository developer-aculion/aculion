r"""
Real-Time AI Vehicle Detection with Zero-Recursion Scrcpy Capture,
Advanced Traffic Analytics & Automated Camera Snapshots (Supabase Partitioned by stat_date)
========================================================================================
Target Window: scrcpy ("SM-A146B")
YOLO Model: runs/detect/runs_vehicle_detect/exp1/weights/best.pt

Features:
1. IST Timezone (Asia/Kolkata) partition by stat_date (YYYY-MM-DD).
2. Startup baseline fetch: resumes from today's existing totals if model restarts.
3. Upsert on (billboard_code, stat_date) matching Supabase schema.
4. Rolling 5-minute flow rate per minute.
5. Automatic midnight rollover.
6. Clean exit handling setting is_live=False.
"""

import sys
import time
import json
import os
from datetime import datetime, timezone, timedelta
from collections import defaultdict, Counter, deque
from pathlib import Path
import numpy as np
import cv2
import win32gui
import win32ui
import win32con
import mss
import pygetwindow as gw
from ultralytics import YOLO
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, Alignment, PatternFill
from openpyxl.utils import get_column_letter
from supabase import create_client, Client
from dotenv import load_dotenv

# ── Load Environment Variables ──
_SCRIPT_DIR = Path(__file__).resolve().parent
_PLATFORM_SERVICES_ENV = _SCRIPT_DIR.parent / ".env"
_ROOT_ENV = Path("e:/Aculion/.env")

if _PLATFORM_SERVICES_ENV.exists():
    load_dotenv(dotenv_path=_PLATFORM_SERVICES_ENV)
elif _ROOT_ENV.exists():
    load_dotenv(dotenv_path=_ROOT_ENV)
else:
    load_dotenv()

# IST Timezone Helper (UTC + 5:30)
IST_TZ = timezone(timedelta(hours=5, minutes=30))

def get_ist_now():
    """Returns the current timezone-aware datetime in IST."""
    return datetime.now(IST_TZ)

def get_today_ist():
    """Returns today's date in IST as YYYY-MM-DD string."""
    return get_ist_now().strftime("%Y-%m-%d")

def get_current_hour_ist():
    """Returns the current hour in IST (0-23)."""
    return get_ist_now().hour

# ==============================================================================
# CONFIGURATION & CONSTANTS
# ==============================================================================
PRIMARY_WINDOW_TITLE = "SM-A146B"
FALLBACK_TITLE_KEYWORDS = ["scrcpy", "SM-", "Samsung", "Android"]

# Model auto-resolving
MODEL_PATH_OVERRIDE = os.environ.get("MODEL_PATH", "")

_MODEL_SEARCH_CANDIDATES = [
    Path("E:/acilion/Aculion 12 (200)/runs/detect/runs_vehicle_detect/exp1/weights/best.pt"),
    Path("E:/Aculion/runs/detect/runs_vehicle_detect/exp1/weights/best.pt"),
    _SCRIPT_DIR / "runs" / "detect" / "runs_vehicle_detect" / "exp1" / "weights" / "best.pt",
    _SCRIPT_DIR / "weights" / "best.pt",
    _SCRIPT_DIR / "best.pt",
]

def resolve_model_path():
    if MODEL_PATH_OVERRIDE:
        return Path(MODEL_PATH_OVERRIDE)
    for candidate in _MODEL_SEARCH_CANDIDATES:
        if candidate.exists():
            return candidate
    for pattern in ("best.pt", "last.pt"):
        matches = list(_SCRIPT_DIR.rglob(pattern))
        if matches:
            return matches[0]
    return _MODEL_SEARCH_CANDIDATES[0]

MODEL_PATH = resolve_model_path()
CONF_THRESHOLD = 0.25
JSON_OUTPUT_FILE = "vehicle_counts.json"
TRACKER_CONFIG = "custom_bytetrack.yaml"

REQUIRED_CLASSES = ["bike", "commercial", "economic", "luxury", "premium", "ultra_luxury", "pedestrian"]
MIN_FRAMES_FOR_VALID_TRACK = 3
EXPOSURE_TIME_MULTIPLIER = 4.0

OCCUPANCY_MULTIPLIERS = {
    "bike": 1.2,
    "commercial": 1.5,
    "economic": 2.0,
    "premium": 2.0,
    "luxury": 1.8,
    "ultra_luxury": 1.8,
    "pedestrian": 1.0,
}

SNAPSHOT_INTERVAL_SEC = 1800.0
SNAPSHOT_DIR = Path("camera_snapshots")
EXCEL_OUTPUT_FILE = "vehicle_counts_live.xlsx"
EXCEL_HISTORY_FILE = "vehicle_counts_history.xlsx"
EXCEL_UPDATE_INTERVAL_SEC = 3.0
SUPABASE_UPDATE_INTERVAL_SEC = 5.0

SUPABASE_URL = os.environ.get("SUPABASE_URL") or os.environ.get("VITE_SUPABASE_URL") or "https://buqtshfptmqieaqcghfx.supabase.co"
SUPABASE_KEY = (
    os.environ.get("SUPABASE_SERVICE_ROLE_KEY")
    or os.environ.get("SUPABASE_KEY")
    or os.environ.get("VITE_SUPABASE_ANON_KEY")
    or "eyJhbGciOiJIUzI1NiIsInR5cCI6IkpXVCJ9.eyJpc3MiOiJzdXBhYmFzZSIsInJlZiI6ImJ1cXRzaGZwdG1xaWVhcWNnaGZ4Iiwicm9sZSI6InNlcnZpY2Vfcm9sZSIsImlhdCI6MTc4MzkwOTYyMiwiZXhwIjoyMDk5NDg1NjIyfQ.f12uC9oK_BzLzlXgy_5ybUAgdHJTY6N7E5VWXXmgr5Q"
)
SUPABASE_TABLE = "traffic_overview"

BILLBOARD_CODE = os.environ.get("BILLBOARD_CODE") or "ACU-BB-0001"
BILLBOARD_NAME = os.environ.get("BILLBOARD_NAME") or "Testing Billboard-1"
CAMERA_FF_CODE = os.environ.get("CAMERA_FF_CODE") or "CAM-FF-001"
CAMERA_BF_CODE = os.environ.get("CAMERA_BF_CODE") or "CAM-BF-001"
RADXA_CODE = os.environ.get("RADXA_CODE") or "RADXA-01"

LABEL_MAP = {
    "bike": "BIKE",
    "commercial": "COMMERICAL_VEHICLES",
    "economic": "ECONOMIC",
    "luxury": "LUXURY",
    "premium": "PREMIUM",
    "ultra_luxury": "ULTRA_LUXURY",
    "pedestrian": "PEDESTRIAN",
}

CLASS_COLORS = {
    "bike": (255, 128, 0),
    "commercial": (0, 165, 255),
    "economic": (0, 255, 0),
    "luxury": (255, 0, 255),
    "premium": (0, 255, 255),
    "ultra_luxury": (0, 0, 255),
    "pedestrian": (255, 255, 255)
}
DEFAULT_COLOR = (200, 200, 200)

# ==============================================================================
# SIMPLE IOU TRACKER FALLBACK
# ==============================================================================
class SimpleIoUTracker:
    def __init__(self, iou_thresh=0.3, max_lost_frames=30):
        self.next_id = 1
        self.tracks = {}
        self.iou_thresh = iou_thresh
        self.max_lost_frames = max_lost_frames

    @staticmethod
    def compute_iou(box1, box2):
        x1 = max(box1[0], box2[0])
        y1 = max(box1[1], box2[1])
        x2 = min(box1[2], box2[2])
        y2 = min(box1[3], box2[3])

        inter_area = max(0, x2 - x1) * max(0, y2 - y1)
        b1_area = (box1[2] - box1[0]) * (box1[3] - box1[1])
        b2_area = (box2[2] - box2[0]) * (box2[3] - box2[1])

        union_area = b1_area + b2_area - inter_area
        return inter_area / union_area if union_area > 0 else 0

    def update(self, boxes):
        assigned_ids = []
        for tid in self.tracks:
            self.tracks[tid]["lost"] += 1

        for box in boxes:
            best_iou = 0
            best_id = None
            for tid, tinfo in self.tracks.items():
                iou = self.compute_iou(box, tinfo["box"])
                if iou > best_iou and iou >= self.iou_thresh:
                    best_iou = iou
                    best_id = tid

            if best_id is not None:
                self.tracks[best_id]["box"] = box
                self.tracks[best_id]["lost"] = 0
                assigned_ids.append(best_id)
            else:
                new_id = self.next_id
                self.next_id += 1
                self.tracks[new_id] = {"box": box, "lost": 0}
                assigned_ids.append(new_id)

        dead_ids = [tid for tid, tinfo in self.tracks.items() if tinfo["lost"] > self.max_lost_frames]
        for tid in dead_ids:
            del self.tracks[tid]

        return assigned_ids

# ==============================================================================
# SCRCPY WINDOW CAPTURER (HWND DEVICE CONTEXT)
# ==============================================================================
class ScrcpyCapturer:
    def __init__(self, target_title=PRIMARY_WINDOW_TITLE):
        self.target_title = target_title
        self.hwnd = None
        self.sct = mss.mss()
        self._find_window()

    def _find_window(self):
        self.hwnd = win32gui.FindWindow(None, self.target_title)
        if not self.hwnd:
            def enum_cb(hwnd, found_list):
                if win32gui.IsWindowVisible(hwnd):
                    title = win32gui.GetWindowText(hwnd)
                    for kw in FALLBACK_TITLE_KEYWORDS:
                        if kw.lower() in title.lower():
                            found_list.append((hwnd, title))
                            break
            found = []
            win32gui.EnumWindows(enum_cb, found)
            if found:
                self.hwnd, title = found[0]
                print(f"[+] Located scrcpy window: '{title}' (HWND: {self.hwnd})")

        if self.hwnd:
            print(f"[+] Successfully bound to scrcpy window (HWND: {self.hwnd})")
            if win32gui.IsIconic(self.hwnd):
                win32gui.ShowWindow(self.hwnd, win32con.SW_RESTORE)
        else:
            print(f"[!] Warning: Could not find scrcpy window matching '{self.target_title}'.")

    def get_client_rect(self):
        if not self.hwnd or not win32gui.IsWindow(self.hwnd):
            self._find_window()
            if not self.hwnd:
                return None
        try:
            pt = win32gui.ClientToScreen(self.hwnd, (0, 0))
            cl_rect = win32gui.GetClientRect(self.hwnd)
            w = cl_rect[2] - cl_rect[0]
            h = cl_rect[3] - cl_rect[1]
            return (pt[0], pt[1], w, h)
        except Exception:
            return None

    def capture_frame(self):
        rect = self.get_client_rect()
        if not rect or rect[2] <= 0 or rect[3] <= 0:
            return None
        x, y, w, h = rect
        try:
            hwnd_dc = win32gui.GetWindowDC(self.hwnd)
            mfc_dc = win32ui.CreateDCFromHandle(hwnd_dc)
            save_dc = mfc_dc.CreateCompatibleDC()

            save_bit_map = win32ui.CreateBitmap()
            save_bit_map.CreateCompatibleBitmap(mfc_dc, w, h)
            save_dc.SelectObject(save_bit_map)

            # PW_RENDERFULLCONTENT = 2
            result = ctypes.windll.user32.PrintWindow(self.hwnd, save_dc.GetSafeHdc(), 2)
            if result == 0:
                result = ctypes.windll.user32.PrintWindow(self.hwnd, save_dc.GetSafeHdc(), 0)

            bmpinfo = save_bit_map.GetInfo()
            bmpstr = save_bit_map.GetBitmapBits(True)
            img = np.frombuffer(bmpstr, dtype=np.uint8).reshape((bmpinfo['bmHeight'], bmpinfo['bmWidth'], 4))

            win32gui.DeleteObject(save_bit_map.GetHandle())
            save_dc.DeleteDC()
            mfc_dc.DeleteDC()
            win32gui.ReleaseDC(self.hwnd, hwnd_dc)

            return cv2.cvtColor(img, cv2.COLOR_BGRA2BGR)
        except Exception:
            try:
                monitor = {"top": y, "left": x, "width": w, "height": h}
                sct_img = self.sct.grab(monitor)
                return cv2.cvtColor(np.array(sct_img), cv2.COLOR_BGRA2BGR)
            except Exception:
                return None

# ==============================================================================
# SUPABASE CLIENT & BASELINE RECOVERY
# ==============================================================================
_supabase_client = None

def get_supabase_client():
    global _supabase_client
    if not SUPABASE_URL or not SUPABASE_KEY:
        return None
    if _supabase_client is None:
        _supabase_client = create_client(SUPABASE_URL, SUPABASE_KEY)
    return _supabase_client

def fetch_today_baseline(billboard_code, stat_date):
    """
    Fetches today's baseline row for this billboard if it already exists in Supabase.
    Allows seamlessly continuing counts upon mid-day script restart.
    """
    client = get_supabase_client()
    if client is None:
        return None
    try:
        res = client.table(SUPABASE_TABLE)\
            .select("*")\
            .eq("billboard_code", billboard_code)\
            .eq("stat_date", stat_date)\
            .limit(1)\
            .maybe_single()\
            .execute()
        if res and res.data:
            print(f"[+] Loaded existing baseline for {billboard_code} on {stat_date}: {res.data.get('total_vehicles', 0)} total vehicles.")
            return res.data
    except Exception as e:
        print(f"[!] Baseline fetch notice: {e}")
    return None

# ==============================================================================
# ANALYTICS ENGINE
# ==============================================================================
def compute_analytics(
    track_class_votes, class_names, track_frame_counts,
    completed_exposure_times, hourly_vehicle_counts,
    start_time, current_time, baseline_data=None,
    crossing_timestamps_queue=None
):
    # Base numbers from today's previous session (if any)
    base_counts = {
        "bike": baseline_data.get("bikes", 0) if baseline_data else 0,
        "commercial": baseline_data.get("commercial", 0) if baseline_data else 0,
        "economic": baseline_data.get("economy", 0) if baseline_data else 0,
        "premium": baseline_data.get("premium", 0) if baseline_data else 0,
        "luxury": baseline_data.get("luxury", 0) if baseline_data else 0,
        "ultra_luxury": baseline_data.get("ultra_luxury", 0) if baseline_data else 0,
        "pedestrian": 0
    }

    # Static vehicle class counts for the current session
    session_counts = defaultdict(int)
    for tid, votes in track_class_votes.items():
        if track_frame_counts.get(tid, 0) >= MIN_FRAMES_FOR_VALID_TRACK:
            best_cid = votes.most_common(1)[0][0]
            cls_name = class_names.get(best_cid, "economic").lower()
            if cls_name in base_counts:
                session_counts[cls_name] += 1
            else:
                session_counts["economic"] += 1

    # Merge baseline with current session
    combined_static_counts = {
        k: base_counts.get(k, 0) + session_counts.get(k, 0)
        for k in REQUIRED_CLASSES
    }
    total_crossed = sum(combined_static_counts[k] for k in REQUIRED_CLASSES if k != "pedestrian")

    # Exposure times
    base_avg_exposure = float(baseline_data.get("avg_exposure_time", 0)) if baseline_data else 0.0
    base_max_exposure = float(baseline_data.get("max_exposure_time", 0)) if baseline_data else 0.0
    
    if len(completed_exposure_times) > 0:
        session_avg_exposure = float(np.mean(completed_exposure_times))
        session_max_exposure = float(np.max(completed_exposure_times))
        if base_avg_exposure > 0:
            avg_exposure_time = (base_avg_exposure + session_avg_exposure) / 2.0
            max_exposure_time = max(base_max_exposure, session_max_exposure)
        else:
            avg_exposure_time = session_avg_exposure
            max_exposure_time = session_max_exposure
    else:
        avg_exposure_time = base_avg_exposure if base_avg_exposure > 0 else 0.0
        max_exposure_time = base_max_exposure if base_max_exposure > 0 else 0.0

    # Peak traffic hour
    peak_traffic_hour = "N/A"
    peak_hour_count = 0
    if hourly_vehicle_counts:
        best_hour = max(hourly_vehicle_counts.items(), key=lambda x: len(x[1]))
        peak_traffic_hour = best_hour[0]
        peak_hour_count = len(best_hour[1])
    elif baseline_data and baseline_data.get("peak_traffic_hour"):
        peak_traffic_hour = baseline_data.get("peak_traffic_hour")

    # Estimated Reach
    estimated_reach = sum(
        combined_static_counts.get(cname, 0) * OCCUPANCY_MULTIPLIERS.get(cname, 1.5)
        for cname in REQUIRED_CLASSES
    )

    # Rolling 5-minute flow rate per minute
    if crossing_timestamps_queue is not None:
        # Purge entries older than 300 seconds (5 min)
        cutoff = current_time - 300.0
        while crossing_timestamps_queue and crossing_timestamps_queue[0] < cutoff:
            crossing_timestamps_queue.popleft()
        
        window_sec = min(max(current_time - start_time, 1.0), 300.0)
        window_min = window_sec / 60.0
        recent_count = len(crossing_timestamps_queue)
        flow_rate_per_min = recent_count / window_min if window_min > 0 else 0.0
    else:
        elapsed_sec = max(current_time - start_time, 1.0)
        elapsed_min = elapsed_sec / 60.0
        flow_rate_per_min = total_crossed / elapsed_min if elapsed_min > 0 else 0.0

    elapsed_sec = max(current_time - start_time, 0.001)

    return {
        "static_counts": combined_static_counts,
        "total_vehicles_crossed": total_crossed,
        "average_exposure_time_sec": round(avg_exposure_time, 2),
        "max_exposure_time_sec": round(max_exposure_time, 2),
        "completed_exposure_sample_size": len(completed_exposure_times),
        "peak_traffic_hour": peak_traffic_hour,
        "peak_hour_count": peak_hour_count,
        "estimated_reach_persons": int(round(estimated_reach)),
        "flow_rate_per_min": round(flow_rate_per_min, 1),
        "elapsed_sec": round(elapsed_sec, 1)
    }

# ==============================================================================
# SUPABASE PUSH FUNCTIONS
# ==============================================================================
def build_supabase_payload(analytics, is_live, stat_date_str):
    sc = analytics["static_counts"]
    ist_now = get_ist_now()
    return {
        "radxa_code": RADXA_CODE,
        "billboard_code": BILLBOARD_CODE,
        "billboard_name": BILLBOARD_NAME,
        "camera_ff_code": CAMERA_FF_CODE,
        "camera_bf_code": CAMERA_BF_CODE,
        "stat_date": stat_date_str,
        "hour": ist_now.hour,
        "total_vehicles": analytics["total_vehicles_crossed"],
        "bikes": sc.get("bike", 0),
        "economy": sc.get("economic", 0),
        "premium": sc.get("premium", 0),
        "luxury": sc.get("luxury", 0),
        "ultra_luxury": sc.get("ultra_luxury", 0),
        "commercial": sc.get("commercial", 0),
        "avg_exposure_time": analytics["average_exposure_time_sec"],
        "max_exposure_time": analytics["max_exposure_time_sec"],
        "estimated_reach": analytics["estimated_reach_persons"],
        "flow_rate": analytics["flow_rate_per_min"],
        "peak_traffic_hour": analytics["peak_traffic_hour"],
        "is_live": is_live,
        "is_legacy": False,
        "last_updated": ist_now.isoformat(),
    }

def push_to_supabase(analytics, is_live=True, stat_date_str=None):
    client = get_supabase_client()
    if client is None:
        return

    if not stat_date_str:
        stat_date_str = get_today_ist()

    payload = build_supabase_payload(analytics, is_live, stat_date_str)

    try:
        client.table(SUPABASE_TABLE).upsert(payload, on_conflict="billboard_code,stat_date").execute()
    except Exception as e:
        print(f"[!] Supabase live push failed: {e}")

def push_to_supabase_history(analytics, stat_date_str=None):
    client = get_supabase_client()
    if client is None:
        return

    if not stat_date_str:
        stat_date_str = get_today_ist()

    payload = build_supabase_payload(analytics, True, stat_date_str)
    history_payload = {
        "radxa_code": payload["radxa_code"],
        "billboard_code": payload["billboard_code"],
        "billboard_name": payload["billboard_name"],
        "camera_ff_code": payload["camera_ff_code"],
        "camera_bf_code": payload["camera_bf_code"],
        "stat_date": stat_date_str,
        "total_vehicles": payload["total_vehicles"],
        "bikes": payload["bikes"],
        "economy": payload["economy"],
        "premium": payload["premium"],
        "luxury": payload["luxury"],
        "ultra_luxury": payload["ultra_luxury"],
        "commercial": payload["commercial"],
        "avg_exposure_time": payload["avg_exposure_time"],
        "max_exposure_time": payload["max_exposure_time"],
        "estimated_reach": payload["estimated_reach"],
        "flow_rate": payload["flow_rate"],
        "peak_traffic_hour": payload["peak_traffic_hour"],
        "recorded_at": get_ist_now().isoformat(),
        "is_live": True
    }

    try:
        client.table("traffic_overview_history").insert(history_payload).execute()
    except Exception as e:
        print(f"[!] History snapshot failed: {e}")

def save_camera_snapshot(frame, snapshot_counter):
    try:
        SNAPSHOT_DIR.mkdir(parents=True, exist_ok=True)
        timestamp_str = get_ist_now().strftime("%Y%m%d_%H%M%S")
        filename = SNAPSHOT_DIR / f"snapshot_{timestamp_str}_{snapshot_counter:03d}.jpg"
        cv2.imwrite(str(filename), frame)
        print(f"[+] Camera Feed Snapshot #{snapshot_counter} saved: '{filename}'")
        return filename
    except Exception as e:
        print(f"[!] Failed to save camera snapshot: {e}")
        return None

# ==============================================================================
# MAIN EXECUTION LOOP
# ==============================================================================
def main():
    print("=" * 65)
    print("  ACULION AI TRAFFIC INTELLIGENCE ENGINE (PARTITIONED BY stat_date)")
    print("=" * 65)

    current_stat_date = get_today_ist()
    print(f"[+] Active Billboard: {BILLBOARD_NAME} ({BILLBOARD_CODE})")
    print(f"[+] Tracking Partition Date (IST): {current_stat_date}")

    # 1. Fetch startup baseline
    baseline_data = fetch_today_baseline(BILLBOARD_CODE, current_stat_date)

    # 2. Load YOLO Model
    print(f"\n[+] Loading YOLO model from: '{MODEL_PATH}'")
    try:
        model = YOLO(str(MODEL_PATH))
        class_names = model.names
        print(f"[+] Model loaded successfully ({len(class_names)} classes).")
    except Exception as e:
        print(f"[ERROR] Failed to load YOLO model: {e}")
        sys.exit(1)

    # 3. Ensure ByteTrack config
    if not Path(TRACKER_CONFIG).exists():
        with open(TRACKER_CONFIG, "w") as f:
            f.write("""tracker_type: bytetrack
track_high_thresh: 0.5
track_low_thresh: 0.1
new_track_thresh: 0.6
track_buffer: 60
match_thresh: 0.8
fuse_score: True
""")

    # 4. Capturer & Tracker setup
    capturer = ScrcpyCapturer(PRIMARY_WINDOW_TITLE)
    fallback_tracker = SimpleIoUTracker()

    window_name = "CCTV Traffic Analytics & Vehicle Detection (scrcpy Feed)"
    cv2.namedWindow(window_name, cv2.WINDOW_NORMAL)
    cv2.setWindowProperty(window_name, cv2.WND_PROP_FULLSCREEN, cv2.WINDOW_FULLSCREEN)

    is_fullscreen = True
    show_hud = True

    # State variables
    track_class_votes = defaultdict(Counter)
    track_timestamps = {}
    track_frame_counts = defaultdict(int)
    hourly_vehicle_counts = defaultdict(set)
    completed_exposure_times = []
    finalized_track_ids = set()
    prev_active_ids = set()
    crossing_timestamps_queue = deque()

    fps = 0.0
    start_time = time.time()
    last_supabase_update = 0.0
    last_snapshot_time = 0.0
    snapshot_counter = 0

    try:
        while True:
            t_loop_start = time.time()

            # Check for midnight rollover in IST
            today_check = get_today_ist()
            if today_check != current_stat_date:
                print(f"[!] Midnight rollover detected ({current_stat_date} -> {today_check}). Archiving previous day partition...")
                # Push final for previous day
                final_analytics = compute_analytics(
                    track_class_votes, class_names, track_frame_counts,
                    completed_exposure_times, hourly_vehicle_counts,
                    start_time, t_loop_start, baseline_data, crossing_timestamps_queue
                )
                push_to_supabase(final_analytics, is_live=False, stat_date_str=current_stat_date)
                
                # Reset state for new day
                current_stat_date = today_check
                baseline_data = None
                track_class_votes.clear()
                track_timestamps.clear()
                track_frame_counts.clear()
                hourly_vehicle_counts.clear()
                completed_exposure_times.clear()
                finalized_track_ids.clear()
                prev_active_ids.clear()
                crossing_timestamps_queue.clear()
                start_time = time.time()
                print(f"[+] Started new partition for {current_stat_date}.")

            # Capture frame
            frame = capturer.capture_frame()
            if frame is None:
                waiting_frame = np.zeros((720, 1280, 3), dtype=np.uint8)
                msg1 = "Waiting for scrcpy window ('SM-A146B')..."
                msg2 = "Ensure scrcpy is open and connected to device."
                cv2.putText(waiting_frame, msg1, (100, 320), cv2.FONT_HERSHEY_SIMPLEX, 1.0, (0, 165, 255), 2)
                cv2.putText(waiting_frame, msg2, (100, 380), cv2.FONT_HERSHEY_SIMPLEX, 0.7, (200, 200, 200), 1)
                cv2.imshow(window_name, waiting_frame)
                key = cv2.waitKey(200) & 0xFF
                if key in (27, ord('q'), ord('Q')):
                    break
                continue

            # Snapshots
            if (t_loop_start - last_snapshot_time) >= SNAPSHOT_INTERVAL_SEC:
                snapshot_counter += 1
                save_camera_snapshot(frame, snapshot_counter)
                last_snapshot_time = t_loop_start

            annotated_frame = frame.copy()

            # YOLO inference
            try:
                results = model.track(source=frame, conf=CONF_THRESHOLD, persist=True, tracker=TRACKER_CONFIG, verbose=False)
                r = results[0] if len(results) > 0 else None
                boxes = r.boxes if r else None
                use_bytetrack = (boxes is not None and boxes.id is not None)
            except Exception:
                use_bytetrack = False
                results = model(frame, conf=CONF_THRESHOLD, verbose=False)
                r = results[0] if len(results) > 0 else None
                boxes = r.boxes if r else None

            current_hour_key = get_ist_now().strftime("%H:00 - %H:59")
            current_active_ids = set()

            if boxes is not None and len(boxes) > 0:
                cls_ids = boxes.cls.int().tolist()
                confs = boxes.conf.tolist()
                xyxy = boxes.xyxy.tolist()
                track_ids = boxes.id.int().tolist() if use_bytetrack else fallback_tracker.update(xyxy)

                for cid, tid, conf, box in zip(cls_ids, track_ids, confs, xyxy):
                    track_class_votes[tid][cid] += 1
                    track_frame_counts[tid] += 1
                    current_active_ids.add(tid)

                    if tid not in track_timestamps:
                        track_timestamps[tid] = {"first_seen": t_loop_start, "last_seen": t_loop_start}
                    else:
                        track_timestamps[tid]["last_seen"] = t_loop_start

                    # Check if this frame confirms a new valid vehicle crossing
                    if track_frame_counts[tid] == MIN_FRAMES_FOR_VALID_TRACK:
                        crossing_timestamps_queue.append(t_loop_start)

                    hourly_vehicle_counts[current_hour_key].add(tid)

                    x1, y1, x2, y2 = map(int, box)
                    cls_name = class_names.get(cid, f"Class_{cid}")
                    color = CLASS_COLORS.get(cls_name.lower(), DEFAULT_COLOR)

                    cv2.rectangle(annotated_frame, (x1, y1), (x2, y2), color, 2)
                    label = f"#{tid} {cls_name} {conf * 100:.0f}%"
                    (tw, th), _ = cv2.getTextSize(label, cv2.FONT_HERSHEY_SIMPLEX, 0.55, 2)
                    cv2.rectangle(annotated_frame, (x1, max(y1 - th - 6, 0)), (x1 + tw + 6, y1), color, -1)
                    text_color = (0, 0, 0) if sum(color) > 400 else (255, 255, 255)
                    cv2.putText(annotated_frame, label, (x1 + 3, max(y1 - 3, th)), cv2.FONT_HERSHEY_SIMPLEX, 0.55, text_color, 2, cv2.LINE_AA)

            # Exposure finalization
            disappeared_ids = prev_active_ids - current_active_ids
            for tid in disappeared_ids:
                if tid in finalized_track_ids:
                    continue
                if track_frame_counts.get(tid, 0) >= MIN_FRAMES_FOR_VALID_TRACK:
                    raw_exposure = track_timestamps[tid]["last_seen"] - track_timestamps[tid]["first_seen"]
                    completed_exposure_times.append(raw_exposure * EXPOSURE_TIME_MULTIPLIER)
                finalized_track_ids.add(tid)
            prev_active_ids = current_active_ids

            # Compute analytics
            analytics = compute_analytics(
                track_class_votes, class_names, track_frame_counts,
                completed_exposure_times, hourly_vehicle_counts,
                start_time, t_loop_start, baseline_data, crossing_timestamps_queue
            )

            # FPS calculation
            t_loop_end = time.time()
            dt = t_loop_end - t_loop_start
            current_fps = 1.0 / dt if dt > 0 else 30.0
            fps = 0.9 * fps + 0.1 * current_fps if fps > 0 else current_fps

            # Throttled Supabase push
            if t_loop_end - last_supabase_update >= SUPABASE_UPDATE_INTERVAL_SEC:
                push_to_supabase(analytics, is_live=True, stat_date_str=current_stat_date)
                push_to_supabase_history(analytics, stat_date_str=current_stat_date)
                last_supabase_update = t_loop_end

            # HUD
            if show_hud:
                hud_w, hud_h = 390, 420
                overlay = annotated_frame.copy()
                cv2.rectangle(overlay, (10, 10), (10 + hud_w, 10 + hud_h), (15, 15, 15), -1)
                cv2.addWeighted(overlay, 0.75, annotated_frame, 0.25, 0, annotated_frame)
                cv2.rectangle(annotated_frame, (10, 10), (10 + hud_w, 10 + hud_h), (0, 255, 0), 1)

                cv2.putText(annotated_frame, "TRAFFIC ANALYTICS HUD", (20, 34), cv2.FONT_HERSHEY_SIMPLEX, 0.65, (0, 255, 255), 2)
                cv2.putText(annotated_frame, f"FPS: {fps:.1f}", (310, 34), cv2.FONT_HERSHEY_SIMPLEX, 0.6, (0, 255, 0), 2)
                cv2.line(annotated_frame, (20, 42), (390, 42), (80, 80, 80), 1)

                y = 65
                cv2.putText(annotated_frame, f"TOTAL VEHICLES CROSSED : {analytics['total_vehicles_crossed']}", (20, y), cv2.FONT_HERSHEY_SIMPLEX, 0.52, (255, 255, 255), 2)
                y += 24
                cv2.putText(annotated_frame, f"AVERAGE EXPOSURE TIME    : {analytics['average_exposure_time_sec']}s", (20, y), cv2.FONT_HERSHEY_SIMPLEX, 0.50, (200, 255, 200), 2)
                y += 22
                cv2.putText(annotated_frame, f"MAX EXPOSURE TIME        : {analytics['max_exposure_time_sec']}s", (20, y), cv2.FONT_HERSHEY_SIMPLEX, 0.50, (200, 255, 200), 2)
                y += 22
                cv2.putText(annotated_frame, f"PEAK TRAFFIC HOUR     : {analytics['peak_traffic_hour']}", (20, y), cv2.FONT_HERSHEY_SIMPLEX, 0.50, (255, 200, 100), 2)
                y += 22
                cv2.putText(annotated_frame, f"ESTIMATED REACH       : {analytics['estimated_reach_persons']} persons", (20, y), cv2.FONT_HERSHEY_SIMPLEX, 0.50, (255, 180, 255), 2)
                y += 22
                cv2.putText(annotated_frame, f"FLOW RATE (Veh/min)   : {analytics['flow_rate_per_min']} v/m", (20, y), cv2.FONT_HERSHEY_SIMPLEX, 0.50, (100, 255, 255), 2)

                y += 22
                sc = analytics["static_counts"]
                cv2.line(annotated_frame, (20, y), (390, y), (80, 80, 80), 1)
                y += 20
                cv2.putText(annotated_frame, "VEHICLE CLASSIFICATION COUNTS:", (20, y), cv2.FONT_HERSHEY_SIMPLEX, 0.50, (0, 255, 255), 2)
                y += 22
                for cname in REQUIRED_CLASSES:
                    cnt = sc.get(cname, 0)
                    c_color = CLASS_COLORS.get(cname, DEFAULT_COLOR)
                    cv2.putText(annotated_frame, f"  - {cname.capitalize():<14}: {cnt}", (20, y), cv2.FONT_HERSHEY_SIMPLEX, 0.48, c_color, 2)
                    y += 20

            cv2.imshow(window_name, annotated_frame)
            key = cv2.waitKey(1) & 0xFF
            if key in (27, ord('q'), ord('Q')):
                break
            elif key in (ord('s'), ord('S')):
                snapshot_counter += 1
                save_camera_snapshot(frame, snapshot_counter)
                last_snapshot_time = time.time()
            elif key in (ord('f'), ord('F')):
                is_fullscreen = not is_fullscreen
                prop = cv2.WINDOW_FULLSCREEN if is_fullscreen else cv2.WINDOW_NORMAL
                cv2.setWindowProperty(window_name, cv2.WND_PROP_FULLSCREEN, prop)
            elif key in (ord('h'), ord('H')):
                show_hud = not show_hud

    except KeyboardInterrupt:
        print("\n[+] Keyboard Interrupt. Exiting...")
    finally:
        cv2.destroyAllWindows()
        final_time = time.time()
        for tid in track_timestamps:
            if tid in finalized_track_ids:
                continue
            if track_frame_counts.get(tid, 0) >= MIN_FRAMES_FOR_VALID_TRACK:
                raw_exposure = track_timestamps[tid]["last_seen"] - track_timestamps[tid]["first_seen"]
                completed_exposure_times.append(raw_exposure * EXPOSURE_TIME_MULTIPLIER)
            finalized_track_ids.add(tid)

        final_analytics = compute_analytics(
            track_class_votes, class_names, track_frame_counts,
            completed_exposure_times, hourly_vehicle_counts,
            start_time, final_time, baseline_data, crossing_timestamps_queue
        )
        push_to_supabase(final_analytics, is_live=False, stat_date_str=current_stat_date)
        print("[+] Final session state saved (is_live=False).")

if __name__ == "__main__":
    main()
